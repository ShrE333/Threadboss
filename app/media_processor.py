from __future__ import annotations

import asyncio
import base64
import io
import logging
import mimetypes
import os
import tempfile
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse, urlunparse

import httpx
import pytesseract
from PIL import Image
from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from .config import Settings
from .models import MediaRef

logger = logging.getLogger(__name__)


class MediaProcessingError(RuntimeError):
    pass


class MediaProcessor:
    def __init__(self, settings: Settings):
        self.settings = settings
        self._whisper_model = None

    def _rewrite_media_url(self, url: str) -> str:
        """WAHA sometimes emits localhost media URLs. Rewrite them to WAHA_BASE_URL."""
        parsed = urlparse(url)
        if parsed.hostname not in {"localhost", "127.0.0.1", "0.0.0.0"}:
            return url
        base = urlparse(self.settings.waha_base_url)
        return urlunparse((base.scheme, base.netloc, parsed.path, parsed.params, parsed.query, parsed.fragment))

    async def download(self, media: MediaRef) -> tuple[bytes, str, str]:
        if not media.url:
            raise MediaProcessingError("Media URL is missing")

        url = self._rewrite_media_url(media.url)
        headers = {"Accept": "*/*"}
        if self.settings.waha_api_key:
            headers["X-Api-Key"] = self.settings.waha_api_key

        try:
            async with httpx.AsyncClient(timeout=self.settings.media_download_timeout_seconds, follow_redirects=True) as client:
                response = await client.get(url, headers=headers)
                response.raise_for_status()
                data = response.content
        except Exception as exc:
            raise MediaProcessingError(f"Could not download media: {exc}") from exc

        if len(data) > self.settings.max_media_bytes:
            raise MediaProcessingError(
                f"File is too large ({len(data) / 1024 / 1024:.1f} MB). "
                f"Limit is {self.settings.max_media_mb} MB."
            )

        mimetype = media.mimetype or response.headers.get("content-type", "").split(";")[0]
        filename = media.filename or Path(urlparse(url).path).name or "attachment"
        if not mimetype:
            mimetype = mimetypes.guess_type(filename)[0] or "application/octet-stream"
        return data, mimetype, filename

    async def transcribe(self, data: bytes, filename: str = "voice.ogg") -> str:
        if not self.settings.enable_local_stt:
            raise MediaProcessingError("Local speech-to-text is disabled")
        return await asyncio.to_thread(self._transcribe_sync, data, filename)

    def _transcribe_sync(self, data: bytes, filename: str) -> str:
        try:
            from faster_whisper import WhisperModel
        except Exception as exc:
            raise MediaProcessingError(f"faster-whisper is unavailable: {exc}") from exc

        if self._whisper_model is None:
            logger.info("Loading faster-whisper model: %s", self.settings.whisper_model)
            self._whisper_model = WhisperModel(
                self.settings.whisper_model,
                device=self.settings.whisper_device,
                compute_type=self.settings.whisper_compute_type,
                download_root=self.settings.whisper_cache_dir,
            )

        suffix = Path(filename).suffix or ".ogg"
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temp:
            temp.write(data)
            temp_path = temp.name

        try:
            segments, _ = self._whisper_model.transcribe(
                temp_path,
                vad_filter=True,
                beam_size=1,
            )
            text = " ".join(segment.text.strip() for segment in segments if segment.text.strip()).strip()
            return text
        except Exception as exc:
            raise MediaProcessingError(f"Speech transcription failed: {exc}") from exc
        finally:
            try:
                os.unlink(temp_path)
            except OSError:
                pass

    async def extract_text(self, data: bytes, mimetype: str, filename: str) -> str:
        return await asyncio.to_thread(self._extract_text_sync, data, mimetype, filename)

    def _extract_text_sync(self, data: bytes, mimetype: str, filename: str) -> str:
        lower_name = filename.lower()
        try:
            if mimetype.startswith("image/"):
                image = Image.open(io.BytesIO(data)).convert("RGB")
                text = pytesseract.image_to_string(image)
                return text.strip()

            if mimetype == "application/pdf" or lower_name.endswith(".pdf"):
                reader = PdfReader(io.BytesIO(data))
                chunks = []
                for page in reader.pages[: self.settings.max_document_pages]:
                    chunks.append(page.extract_text() or "")
                return "\n".join(chunks).strip()

            if mimetype in {
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "application/msword",
            } or lower_name.endswith(".docx"):
                doc = Document(io.BytesIO(data))
                return "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()

            if mimetype == "application/vnd.openxmlformats-officedocument.presentationml.presentation" or lower_name.endswith(".pptx"):
                prs = Presentation(io.BytesIO(data))
                chunks = []
                for slide in prs.slides[: self.settings.max_document_pages]:
                    for shape in slide.shapes:
                        text = getattr(shape, "text", "")
                        if text and text.strip():
                            chunks.append(text.strip())
                return "\n".join(chunks).strip()

            if mimetype.startswith("text/") or lower_name.endswith((".txt", ".md", ".csv", ".json")):
                return data.decode("utf-8", errors="replace").strip()

            return ""
        except Exception as exc:
            raise MediaProcessingError(f"Could not extract text from {filename}: {exc}") from exc

    async def vision_describe(self, data: bytes, mimetype: str, prompt: str) -> Optional[str]:
        """Optional Gemini vision. OCR is used when this is not configured."""
        if not self.settings.gemini_api_key or not self.settings.vision_model:
            return None

        url = (
            "https://generativelanguage.googleapis.com/v1beta/models/"
            f"{self.settings.vision_model}:generateContent"
        )
        payload = {
            "contents": [
                {
                    "parts": [
                        {"text": prompt or "Describe this image and extract the important information."},
                        {
                            "inlineData": {
                                "mimeType": mimetype or "image/jpeg",
                                "data": base64.b64encode(data).decode("ascii"),
                            }
                        },
                    ]
                }
            ]
        }
        try:
            async with httpx.AsyncClient(timeout=self.settings.vision_timeout_seconds) as client:
                response = await client.post(url, params={"key": self.settings.gemini_api_key}, json=payload)
                response.raise_for_status()
                result = response.json()
            candidates = result.get("candidates") or []
            if not candidates:
                return None
            parts = (((candidates[0] or {}).get("content") or {}).get("parts") or [])
            text = "\n".join(str(p.get("text", "")).strip() for p in parts if p.get("text"))
            return text.strip() or None
        except Exception as exc:
            logger.warning("Gemini vision failed, falling back to OCR: %s", exc)
            return None

    async def context_for_memory(self, media: MediaRef) -> str:
        """Local-first extraction for passive knowledge indexing.

        Unlike self-chat vision, this avoids spending Gemini requests on every
        image received in normal chats. Images/docs use OCR/text extraction;
        audio uses local Whisper only when enabled by settings.
        """
        data, mimetype, filename = await self.download(media)
        lower = filename.lower()
        if mimetype.startswith("audio/") or lower.endswith((".ogg", ".opus", ".mp3", ".wav", ".m4a", ".aac")):
            if not self.settings.index_normal_chat_audio:
                return ""
            transcript = await self.transcribe(data, filename)
            return f"Voice-note transcript:\n{transcript}".strip()

        if mimetype.startswith("image/"):
            if not self.settings.index_normal_chat_images:
                return ""
            text = await self.extract_text(data, mimetype, filename)
            return f"Image OCR text:\n{text}" if text else ""

        if not self.settings.index_normal_chat_documents:
            return ""
        text = await self.extract_text(data, mimetype, filename)
        if text:
            if len(text) > self.settings.media_context_max_chars:
                text = text[: self.settings.media_context_max_chars] + "\n[truncated]"
            return f"Attachment '{filename}' extracted text:\n{text}"
        return ""

    async def context_for_media(self, media: MediaRef, user_prompt: str = "") -> str:
        data, mimetype, filename = await self.download(media)

        if mimetype.startswith("audio/") or filename.lower().endswith((".ogg", ".opus", ".mp3", ".wav", ".m4a", ".aac")):
            transcript = await self.transcribe(data, filename)
            return f"Voice-note transcript:\n{transcript}".strip()

        if mimetype.startswith("image/"):
            visual = await self.vision_describe(data, mimetype, user_prompt)
            if visual:
                return f"Image analysis:\n{visual}"
            text = await self.extract_text(data, mimetype, filename)
            return f"Image OCR text:\n{text}" if text else "Image received, but no readable text was extracted."

        text = await self.extract_text(data, mimetype, filename)
        if text:
            if len(text) > self.settings.media_context_max_chars:
                text = text[: self.settings.media_context_max_chars] + "\n[truncated]"
            return f"Attachment '{filename}' extracted text:\n{text}"

        return f"Attachment received: {filename} ({mimetype})"
